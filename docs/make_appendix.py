"""
Builds docs/SpeakSafe_Appendix.docx in the same format as the reference report's appendix:
FILE STRUCTURE, then MODULE 1, MODULE 2, ... with the source code, then SAMPLE OUTPUT.

Run from the project root after changing the code:   python3 docs/make_appendix.py
Needs:  pip install python-docx
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parent.parent
JAVA = "backend/src/main/java/com/speaksafe/"
RES = "backend/src/main/resources/"

MODULES = [
    ("Data Models", [JAVA + "model/Category.java", JAVA + "model/Action.java",
                     JAVA + "model/AnalysisResult.java", JAVA + "model/Report.java"]),
    ("Preprocessing Module", [JAVA + "nlp/TextPreprocessor.java"]),
    ("Feature Extraction Module (TF-IDF)", [JAVA + "nlp/TfIdfVectorizer.java"]),
    ("Classification Module", [JAVA + "ml/Classifier.java", JAVA + "ml/NaiveBayesClassifier.java",
                               JAVA + "ml/SvmClassifier.java"]),
    ("Training and Evaluation", [JAVA + "ml/DatasetLoader.java", JAVA + "ml/DetectionPipeline.java",
                                 JAVA + "ml/ModelEvaluator.java"]),
    ("Categorization, Scoring and Action Module", [JAVA + "service/ToxicityDetector.java"]),
    ("Database Module", [RES + "schema.sql", JAVA + "service/ReportRepository.java"]),
    ("Input Module (REST API)", [JAVA + "controller/DetectionController.java"]),
    ("Main Application", [JAVA + "SpeakSafeApplication.java", RES + "application.properties"]),
    ("Browser Extension", ["extension/manifest.json", "extension/background.js", "extension/content.js",
                           "extension/content.css", "extension/popup.html", "extension/popup.js"]),
    ("Web Interface (Message Checker and Moderator Dashboard)", [RES + "static/index.html",
                                                                  RES + "static/dashboard.html"]),
]

TREE = """SpeakSafe
├── backend
│   ├── pom.xml
│   └── src/main
│       ├── java/com/speaksafe
│       │   ├── SpeakSafeApplication.java
│       │   ├── model
│       │   │   ├── Category.java
│       │   │   ├── Action.java
│       │   │   ├── AnalysisResult.java
│       │   │   └── Report.java
│       │   ├── nlp
│       │   │   ├── TextPreprocessor.java
│       │   │   └── TfIdfVectorizer.java
│       │   ├── ml
│       │   │   ├── Classifier.java
│       │   │   ├── NaiveBayesClassifier.java
│       │   │   ├── SvmClassifier.java
│       │   │   ├── DatasetLoader.java
│       │   │   ├── DetectionPipeline.java
│       │   │   └── ModelEvaluator.java
│       │   ├── service
│       │   │   ├── ToxicityDetector.java
│       │   │   └── ReportRepository.java
│       │   └── controller
│       │       └── DetectionController.java
│       └── resources
│           ├── application.properties
│           ├── schema.sql
│           ├── dataset/speaksafe_dataset.csv
│           └── static (index.html, dashboard.html, style.css)
└── extension
    ├── manifest.json
    ├── background.js
    ├── content.js
    ├── content.css
    ├── popup.html
    ├── popup.js
    ├── popup.css
    └── icons"""

SCREENSHOTS = [
    ("forum.png", "Browser extension blurring toxic messages in a group chat and warning while typing"),
    ("popup.png", "Extension popup with settings, counters and quick test"),
    ("checker.png", "Web message checker showing toxicity score and category"),
    ("dashboard.png", "Moderator dashboard with flagged messages and model evaluation"),
]


def heading(doc, text, size=14, center=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def body(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"


def code(doc, text):
    for line in text.rstrip().split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line.replace("\t", "    "))
        run.font.name = "Courier New"
        run.font.size = Pt(9)


def main():
    doc = Document()
    heading(doc, "APPENDIX", 16, center=True)
    body(doc, "The appendix includes supplementary information, sample outputs, and additional "
              "details that support the main report.")
    heading(doc, "SOURCE CODE:")
    heading(doc, "FILE STRUCTURE", 12)
    code(doc, TREE)

    for number, (title, files) in enumerate(MODULES, start=1):
        doc.add_page_break()
        heading(doc, f"MODULE {number} – {title.upper()}")
        for file in files:
            heading(doc, Path(file).name, 12)
            code(doc, (ROOT / file).read_text(encoding="utf-8"))
            doc.add_paragraph()

    doc.add_page_break()
    heading(doc, "SAMPLE OUTPUT")
    heading(doc, "Model evaluation (java -jar speaksafe.jar --evaluate)", 12)
    code(doc, (ROOT / "docs/evaluation_output.txt").read_text(encoding="utf-8"))
    for number, (image, caption) in enumerate(SCREENSHOTS, start=1):
        path = ROOT / "docs/screenshots" / image
        if path.exists():
            doc.add_paragraph()
            doc.add_picture(str(path), width=Inches(5.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph(f"Fig A.{number}  {caption}")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    out = ROOT / "docs/SpeakSafe_Appendix.docx"
    doc.save(out)
    print("Saved", out)


if __name__ == "__main__":
    main()
